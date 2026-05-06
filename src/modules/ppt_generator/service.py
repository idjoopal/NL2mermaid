from __future__ import annotations

import io
import yaml
from datetime import datetime
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.modules.shared.types import ChartData, PptContent, SlideContent
from src.utils.logger import get_logger

from .config.config import ppt_generator_config
from .exceptions import SlideGenerationError, TemplateNotFoundError
from .types import PlaceholderSpec, PptResult, SlideDefinition, TemplateDefinition, ThemeDefinition

logger = get_logger("ppt_generator")

CHART_TYPE_MAP = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}


class PptGeneratorService:
    SLIDE_WIDTH = Inches(13.33)
    SLIDE_HEIGHT = Inches(7.5)
    MARGIN = Inches(0.5)
    TITLE_HEIGHT = Inches(1.2)
    CONTENT_TOP = Inches(1.5)
    CONTENT_HEIGHT = Inches(5.6)

    def __init__(self):
        self._templates: dict[str, TemplateDefinition] = {}
        self._load_builtin_templates()

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def get_template(self, name: str) -> TemplateDefinition:
        if name not in self._templates:
            self._try_load_custom_template(name)
        if name not in self._templates:
            raise TemplateNotFoundError(f"템플릿을 찾을 수 없습니다: {name}")
        return self._templates[name]

    def list_templates(self) -> list[str]:
        custom_dir = Path(ppt_generator_config.custom_templates_dir)
        if custom_dir.exists():
            for f in custom_dir.glob("*.yaml"):
                if f.stem not in self._templates:
                    self._try_load_custom_template(f.stem)
        return sorted(self._templates.keys())

    async def execute(
        self,
        content: PptContent,
        template_name: str = "business",
        save_to_file: bool = True,
    ) -> PptResult:
        template = self.get_template(template_name)
        prs = self._create_presentation()
        slide_def_map = {s.id: s for s in template.slides}

        # Reorder content slides to match template definition order
        content_map = {sc.slide_id: sc for sc in content.slides}
        ordered_slides = [
            content_map[sid]
            for sid in (s.id for s in template.slides)
            if sid in content_map
        ]

        for slide_content in ordered_slides:
            slide_def = slide_def_map.get(slide_content.slide_id)
            if not slide_def:
                logger.warning("슬라이드 정의 없음: %s", slide_content.slide_id)
                continue
            try:
                self._add_slide(prs, slide_content, slide_def, template.theme)
            except Exception as e:
                logger.error("슬라이드 생성 오류: slide_id=%s, error=%s", slide_content.slide_id, e)
                raise SlideGenerationError(f"슬라이드 생성 실패 ({slide_content.slide_id}): {e}")

        buf = io.BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        file_path = None
        if save_to_file:
            file_path = self._save_file(pptx_bytes)

        return PptResult(
            pptx_bytes=pptx_bytes,
            file_path=file_path,
            slide_count=len(prs.slides),
            template_name=template_name,
        )

    # ------------------------------------------------------------------
    # Presentation / Slide builders
    # ------------------------------------------------------------------

    def _create_presentation(self) -> Presentation:
        prs = Presentation()
        prs.slide_width = self.SLIDE_WIDTH
        prs.slide_height = self.SLIDE_HEIGHT
        return prs

    def _add_slide(
        self,
        prs: Presentation,
        content: SlideContent,
        slide_def: SlideDefinition,
        theme: ThemeDefinition,
    ) -> None:
        blank = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(blank)

        # White background
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        builders = {
            "title_slide": self._build_title_slide,
            "content_slide": self._build_content_slide,
            "table_slide": self._build_table_slide,
            "chart_slide": self._build_chart_slide,
            "two_column": self._build_two_column_slide,
        }
        build_fn = builders.get(slide_def.type, self._build_content_slide)
        build_fn(slide, content, slide_def, theme)

    # ---- title_slide -------------------------------------------------

    def _build_title_slide(
        self, slide, content: SlideContent, slide_def: SlideDefinition, theme: ThemeDefinition
    ):
        primary = self._rgb(theme.primary_color)
        w, h = self.SLIDE_WIDTH, self.SLIDE_HEIGHT

        # Colored upper block (65% height)
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, w, int(h * 0.65))
        rect.fill.solid()
        rect.fill.fore_color.rgb = primary
        rect.line.fill.background()

        self._textbox(
            slide, content.title,
            self.MARGIN, Inches(2.0), w - 2 * self.MARGIN, Inches(1.5),
            Pt(theme.font_size_title), bold=True,
            color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER,
            font=theme.font_family,
        )
        if content.subtitle:
            self._textbox(
                slide, content.subtitle,
                self.MARGIN, Inches(3.7), w - 2 * self.MARGIN, Inches(0.8),
                Pt(theme.font_size_subtitle), bold=False,
                color=RGBColor(0xCC, 0xD8, 0xFF), align=PP_ALIGN.CENTER,
                font=theme.font_family,
            )

    # ---- content_slide -----------------------------------------------

    def _build_content_slide(
        self, slide, content: SlideContent, slide_def: SlideDefinition, theme: ThemeDefinition
    ):
        self._add_title_bar(slide, content.title, theme)
        body_color = RGBColor(0x33, 0x33, 0x33)

        if content.bullets:
            text = "\n".join(f"•  {b}" for b in content.bullets)
        elif content.body:
            text = content.body
        else:
            return

        self._textbox(
            slide, text,
            self.MARGIN, self.CONTENT_TOP,
            self.SLIDE_WIDTH - 2 * self.MARGIN, self.CONTENT_HEIGHT,
            Pt(theme.font_size_body), bold=False,
            color=body_color, align=PP_ALIGN.LEFT,
            font=theme.font_family, wrap=True,
        )

    # ---- table_slide -------------------------------------------------

    def _build_table_slide(
        self, slide, content: SlideContent, slide_def: SlideDefinition, theme: ThemeDefinition
    ):
        self._add_title_bar(slide, content.title, theme)
        if not content.table_data or len(content.table_data) < 1:
            return

        primary = self._rgb(theme.primary_color)
        secondary = self._rgb(theme.secondary_color)
        rows = len(content.table_data)
        cols = max(len(r) for r in content.table_data)
        w = self.SLIDE_WIDTH - 2 * self.MARGIN

        tbl_shape = slide.shapes.add_table(
            rows, cols,
            self.MARGIN, self.CONTENT_TOP, w, self.CONTENT_HEIGHT,
        )
        tbl = tbl_shape.table

        col_w = int(w / cols)
        for c in range(cols):
            tbl.columns[c].width = col_w

        for r_i, row in enumerate(content.table_data):
            for c_i in range(cols):
                cell_val = row[c_i] if c_i < len(row) else ""
                cell = tbl.cell(r_i, c_i)

                tf = cell.text_frame
                para = tf.paragraphs[0]
                run = para.runs[0] if para.runs else para.add_run()

                run.text = str(cell_val)
                run.font.size = Pt(theme.font_size_body - 2)
                run.font.name = theme.font_family

                if r_i == 0:  # header
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = primary
                else:
                    run.font.bold = False
                    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = (
                        secondary if r_i % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF)
                    )

    # ---- chart_slide -------------------------------------------------

    def _build_chart_slide(
        self, slide, content: SlideContent, slide_def: SlideDefinition, theme: ThemeDefinition
    ):
        self._add_title_bar(slide, content.title, theme)
        if not content.chart_data:
            return

        cd: ChartData = content.chart_data
        chart_type_key = cd.chart_type or slide_def.chart_type or "bar"
        xl_type = CHART_TYPE_MAP.get(chart_type_key, XL_CHART_TYPE.BAR_CLUSTERED)

        chart_data = CategoryChartData()
        chart_data.categories = cd.categories

        # Pie charts only support a single series
        series_list = cd.series[:1] if chart_type_key == "pie" else cd.series
        for series in series_list:
            chart_data.add_series(
                series.get("name", ""),
                tuple(float(v) for v in series.get("values", [])),
            )

        w = self.SLIDE_WIDTH - 2 * self.MARGIN
        slide.shapes.add_chart(
            xl_type,
            self.MARGIN, self.CONTENT_TOP, w, self.CONTENT_HEIGHT,
            chart_data,
        )

    # ---- two_column --------------------------------------------------

    def _build_two_column_slide(
        self, slide, content: SlideContent, slide_def: SlideDefinition, theme: ThemeDefinition
    ):
        self._add_title_bar(slide, content.title, theme)
        body_color = RGBColor(0x33, 0x33, 0x33)
        usable_w = self.SLIDE_WIDTH - 3 * self.MARGIN
        col_w = int(usable_w / 2)
        bullets = content.bullets or []
        mid = len(bullets) // 2

        for i, items in enumerate([bullets[:mid], bullets[mid:]]):
            left = self.MARGIN + i * (col_w + self.MARGIN)
            text = "\n".join(f"•  {b}" for b in items)
            if text:
                self._textbox(
                    slide, text, left, self.CONTENT_TOP, col_w, self.CONTENT_HEIGHT,
                    Pt(theme.font_size_body), False, body_color,
                    PP_ALIGN.LEFT, theme.font_family, wrap=True,
                )

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def _add_title_bar(self, slide, title: str, theme: ThemeDefinition):
        primary = self._rgb(theme.primary_color)
        w = self.SLIDE_WIDTH
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, w, self.TITLE_HEIGHT)
        rect.fill.solid()
        rect.fill.fore_color.rgb = primary
        rect.line.fill.background()
        self._textbox(
            slide, title,
            self.MARGIN, Inches(0.15), w - 2 * self.MARGIN, self.TITLE_HEIGHT,
            Pt(theme.font_size_heading), bold=True,
            color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.LEFT,
            font=theme.font_family,
        )

    def _textbox(
        self, slide, text: str, left, top, width, height,
        size=Pt(18), bold=False, color=None,
        align=PP_ALIGN.LEFT, font="맑은 고딕", wrap=True,
    ):
        txb = slide.shapes.add_textbox(left, top, width, height)
        tf = txb.text_frame
        tf.word_wrap = wrap

        lines = text.split("\n")
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            run = para.add_run()
            run.text = line
            run.font.size = size
            run.font.bold = bold
            run.font.name = font
            if color:
                run.font.color.rgb = color

    @staticmethod
    def _rgb(hex_str: str) -> RGBColor:
        h = hex_str.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    # ------------------------------------------------------------------
    # Template loading
    # ------------------------------------------------------------------

    def _load_builtin_templates(self):
        tpl_dir = Path(__file__).parent / "templates"
        for f in tpl_dir.glob("*.yaml"):
            tpl = self._parse_yaml(f)
            if tpl:
                self._templates[tpl.name] = tpl

    def _try_load_custom_template(self, name: str):
        custom_dir = Path(ppt_generator_config.custom_templates_dir)
        f = custom_dir / f"{name}.yaml"
        if f.exists():
            tpl = self._parse_yaml(f)
            if tpl:
                self._templates[tpl.name] = tpl
                # Also register by filename stem if internal name differs
                if tpl.name != name:
                    self._templates[name] = tpl

    def _parse_yaml(self, path: Path) -> Optional[TemplateDefinition]:
        try:
            with open(path, encoding="utf-8") as f:
                data = yaml.safe_load(f)

            td = data.get("theme", {})
            theme = ThemeDefinition(
                primary_color=td.get("primary_color", "2B4E9E"),
                secondary_color=td.get("secondary_color", "E8EEF8"),
                accent_color=td.get("accent_color", "F0A500"),
                font_family=td.get("font_family", ppt_generator_config.font_fallback),
                font_size_title=td.get("font_size_title", 36),
                font_size_subtitle=td.get("font_size_subtitle", 24),
                font_size_heading=td.get("font_size_heading", 24),
                font_size_body=td.get("font_size_body", 18),
            )

            slides = [
                SlideDefinition(
                    id=s["id"],
                    type=s["type"],
                    chart_type=s.get("chart_type"),
                    placeholders=[
                        PlaceholderSpec(
                            name=p["name"],
                            type=p.get("type", "text"),
                            required=p.get("required", False),
                        )
                        for p in s.get("placeholders", [])
                    ],
                )
                for s in data.get("slides", [])
            ]

            return TemplateDefinition(
                name=data["name"],
                description=data.get("description", ""),
                slides=slides,
                theme=theme,
            )
        except Exception as e:
            logger.error("템플릿 파싱 오류: path=%s, error=%s", path, e)
            return None

    def _save_file(self, pptx_bytes: bytes) -> str:
        out_dir = Path(ppt_generator_config.output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_path = out_dir / f"{ts}.pptx"
        file_path.write_bytes(pptx_bytes)
        return str(file_path)


ppt_generator_service = PptGeneratorService()
